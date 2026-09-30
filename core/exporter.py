"""Excel export of the full comparison with an audit trail, plus an AI-drafted award memo grounded on
numbers produced by the deterministic award engine."""
import json, time
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from . import store, llm, award
from .config import EXPORT_DIR

FILL = {"high": "E2F0D9", "confirmed": "C6E0B4", "medium": "FFF2CC", "low": "F8CBAD"}

def comparison_matrix():
    q = store.read_sql("SELECT line_id, vendor_id, landed_inr_per_uom, confidence, status FROM quotes")
    return q.pivot(index="line_id", columns="vendor_id", values="landed_inr_per_uom"), q.pivot(index="line_id", columns="vendor_id", values="confidence")

def export_workbook(sc=None, filename=None):
    filename = filename or f"RFx_comparison_{time.strftime('%Y%m%d_%H%M')}.xlsx"
    path = EXPORT_DIR / filename
    lines = store.read_sql("SELECT * FROM lines"); vend = store.read_sql("SELECT * FROM vendors")
    val, conf = comparison_matrix()
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        vend.to_excel(xw, sheet_name="Vendors", index=False)
        m = lines[["line_id", "description", "annual_qty", "uom", "last_year_price_inr"]].merge(val.reset_index(), on="line_id", how="left")
        m.to_excel(xw, sheet_name="Comparison (landed INR)", index=False)
        ws = xw.sheets["Comparison (landed INR)"]
        for ci, vid in enumerate(val.columns, start=6):
            for ri, lid in enumerate(m.line_id, start=2):
                c = conf.loc[lid, vid] if lid in conf.index else None
                if c in FILL:
                    ws.cell(ri, ci).fill = PatternFill("solid", fgColor=FILL[c])
        ws.cell(len(m) + 3, 1, "Cell colour = confidence: green high/confirmed, yellow medium, red low. Blank = not quoted.")
        if sc:
            sc["award"].to_excel(xw, sheet_name="Award scenario", index=False)
            pd.DataFrame([{"k": k, "v": json.dumps(v, default=str) if not isinstance(v, (int, float, str)) else v}
                          for k, v in award.scenario_summary(sc).items() if k != "lines"]).to_excel(xw, sheet_name="Award summary", index=False)
        store.read_sql("SELECT vendor_id, line_id, status, raw_price, raw_unit, raw_currency, price_inr_per_uom, freight_inr_per_uom, landed_inr_per_uom, is_estimate, confidence, review_status, derivation, source_file, source_location, source_snippet, flags, override_note FROM quotes").to_excel(xw, sheet_name="Quote details + sources", index=False)
        store.read_sql("SELECT * FROM flags ORDER BY severity").to_excel(xw, sheet_name="Flags", index=False)
        store.read_sql("SELECT * FROM questionnaire").to_excel(xw, sheet_name="Quality gate", index=False)
        store.read_sql("SELECT * FROM terms").to_excel(xw, sheet_name="Terms", index=False)
        store.read_sql("SELECT * FROM documents").to_excel(xw, sheet_name="Documents", index=False)
        pd.DataFrame(store.events(1000), columns=["ts", "actor", "event", "detail"]).assign(
            ts=lambda d: pd.to_datetime(d.ts, unit="s")).to_excel(xw, sheet_name="Audit log", index=False)
        for ws in xw.sheets.values():
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3B57")
            for i, col in enumerate(ws.columns, 1):
                ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, max(len(str(c.value or "")) for c in list(col)[:50]) + 2))
            ws.freeze_panes = "B2"
    store.log("export", {"file": filename}, actor="buyer")
    return path

MEMO_SYSTEM = """You write award recommendation memos for a procurement head. Use ONLY the numbers in the JSON you are
given - never compute new totals or invent figures. Structure: Recommendation (2-3 sentences), Commercials (table),
Why these vendors (quality gate, coverage, terms), Risks & open items (every caveat, needs-review price, estimate,
uncovered line, conditional vendor), Next steps. Markdown. Under 450 words. Use lakh/crore."""

def award_memo(rfx, sc, extra_context=""):
    summ = award.scenario_summary(sc)
    vend = store.read_sql("SELECT vendor_id, vendor_name, gate_status, gate_reasons, payment_days, validity_days, lines_priced FROM vendors").to_dict("records")
    crit = store.read_sql("SELECT vendor_id, line_id, code, message FROM flags WHERE severity IN ('critical','review') AND resolved=0").head(60).to_dict("records")
    terms = store.read_sql("SELECT vendor_id, term, value, deviation FROM terms WHERE deviation IN ('worse','better','conditional')").to_dict("records")
    prompt = json.dumps({"rfx": {k: rfx.get(k) for k in ("rfx_no", "title", "due", "contract_period", "annual_budget_inr")},
                         "scenario": summ, "vendors": vend, "open_flags": crit, "term_deviations": terms, "buyer_note": extra_context}, default=str)
    return llm.text(MEMO_SYSTEM, prompt, max_tokens=2500)
